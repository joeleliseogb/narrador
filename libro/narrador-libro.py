#!/usr/bin/env python3
"""
Narrador de libros — TXT, EPUB, PDF, Word (.doc/.docx), ODT, RTF (RHVoice / Edge).
Uso:
  narrador-libro.py mi-libro.pdf
  narrador-libro.py mi-libro.epub --engine edge --voice edge:es-ES-ElviraNeural
  narrador-libro.py libro.pdf --from-chapter 12 --rate 0.95
"""

from __future__ import annotations

import argparse
import importlib.util
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
HOST_PY = ROOT / "native" / "narrador-host.py"

MAX_CHUNK = 480


def load_tts():
    if not HOST_PY.is_file():
        sys.exit(f"No se encuentra el motor de voz: {HOST_PY}")
    spec = importlib.util.spec_from_file_location("narrador_host", HOST_PY)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_paragraph(paragraph: str, max_len: int = MAX_CHUNK) -> list[str]:
    paragraph = re.sub(r"\s+", " ", paragraph).strip()
    if not paragraph:
        return []
    if len(paragraph) <= max_len:
        return [paragraph]

    parts = re.split(r"(?<=[.!?…])\s+", paragraph)
    chunks: list[str] = []
    buf = ""
    for part in parts:
        if not part:
            continue
        candidate = f"{buf} {part}".strip() if buf else part
        if len(candidate) <= max_len:
            buf = candidate
        else:
            if buf:
                chunks.append(buf)
            if len(part) <= max_len:
                buf = part
            else:
                for i in range(0, len(part), max_len):
                    chunks.append(part[i : i + max_len])
                buf = ""
    if buf:
        chunks.append(buf)
    return chunks


def read_txt(path: Path) -> list[tuple[str, str]]:
    text = normalize_text(path.read_text(encoding="utf-8", errors="replace"))
    return sections_from_plain_text(text, path.stem)


def sections_from_plain_text(text: str, doc_name: str) -> list[tuple[str, str]]:
    text = normalize_text(text)
    if not text:
        return []
    blocks = [b.strip() for b in re.split(r"\n\s*\n+", text) if b.strip()]
    if not blocks:
        return []
    if len(blocks) == 1:
        return [(doc_name, blocks[0])]

    sections: list[tuple[str, str]] = []
    buf: list[str] = []
    buf_title = doc_name
    part_num = 0

    heading_line = re.compile(
        r"^(?:capítulo|capitulo|chapter|parte|part|sección|seccion|section)\s+[\dIVXLCivxlcáéíóúÁÉÍÓÚ]+",
        re.IGNORECASE,
    )

    def flush_buf() -> None:
        nonlocal part_num, buf, buf_title
        if not buf:
            return
        part_num += 1
        title = buf_title if part_num == 1 and buf_title != doc_name else f"{doc_name} · {part_num}"
        sections.append((title, "\n\n".join(buf)))
        buf = []
        buf_title = doc_name

    for block in blocks:
        first_line = block.split("\n", 1)[0].strip()
        if heading_line.match(first_line) or (
            len(first_line) < 80 and first_line.isupper() and len(first_line.split()) <= 8
        ):
            flush_buf()
            buf_title = first_line
            rest = block[len(first_line) :].strip()
            if rest:
                buf.append(rest)
        else:
            buf.append(block)

    flush_buf()
    return sections if sections else [(doc_name, text)]


def is_docx_heading(style_name: str) -> bool:
    s = (style_name or "").lower().replace("_", " ")
    return (
        s.startswith("heading")
        or s.startswith("título")
        or s.startswith("titulo")
        or s.startswith("title")
        or s.startswith("encabezado")
    )


def read_docx(path: Path) -> list[tuple[str, str]] | None:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        return None

    doc = Document(str(path))
    sections: list[tuple[str, str]] = []
    title = path.stem
    paras: list[str] = []

    def flush() -> None:
        nonlocal paras, title
        if paras:
            sections.append((title, "\n\n".join(paras)))
            paras = []

    def add_paragraph(p: Paragraph) -> None:
        text = (p.text or "").replace("\r", "").strip()
        if not text:
            return
        style = p.style.name if p.style else ""
        if is_docx_heading(style):
            flush()
            title = text
        else:
            paras.append(text)

    for block in doc.element.body:
        tag = block.tag.split("}")[-1] if "}" in block.tag else block.tag
        if tag == "p":
            add_paragraph(Paragraph(block, doc))
        elif tag == "tbl":
            flush()
            table = Table(block, doc)
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paras.append(" — ".join(cells))
            flush()

    flush()
    return sections if sections else None


def find_soffice() -> str | None:
    for name in ("soffice", "libreoffice", "lowriter"):
        found = shutil.which(name)
        if found:
            return found
    return None


def convert_to_txt_with_soffice(path: Path) -> str | None:
    soffice = find_soffice()
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        outdir = Path(tmp)
        try:
            proc = subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--nologo",
                    "--nofirststartwizard",
                    "--convert-to",
                    "txt:Text",
                    "--outdir",
                    str(outdir),
                    str(path.resolve()),
                ],
                capture_output=True,
                text=True,
                timeout=240,
                check=False,
            )
        except (OSError, subprocess.TimeoutExpired):
            return None
        if proc.returncode != 0:
            return None
        candidates = sorted(outdir.glob("*.txt"))
        if not candidates:
            fallback = outdir / f"{path.stem}.txt"
            if fallback.is_file():
                candidates = [fallback]
        if not candidates:
            return None
        return candidates[0].read_text(encoding="utf-8", errors="replace")


def convert_doc_with_antiword(path: Path) -> str | None:
    antiword = shutil.which("antiword")
    if not antiword:
        return None
    try:
        proc = subprocess.run(
            [antiword, "-w", "0", str(path)],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired):
        return None
    if proc.returncode != 0 or not (proc.stdout or "").strip():
        return None
    return proc.stdout


def read_office(path: Path) -> list[tuple[str, str]]:
    suffix = path.suffix.lower()

    if suffix in (".docx", ".docm"):
        try:
            sections = read_docx(path)
            if sections:
                return sections
        except Exception as exc:
            print(f"Aviso: python-docx falló ({exc}), probando LibreOffice…", file=sys.stderr)

    plain: str | None = None
    if suffix == ".doc":
        plain = convert_doc_with_antiword(path)
    if not plain:
        plain = convert_to_txt_with_soffice(path)

    if not plain or not plain.strip():
        sys.exit(
            f"No se pudo extraer texto de {path.name}.\n"
            "  • DOCX: pip install --user python-docx\n"
            "  • DOC/ODT/RTF: sudo dnf install libreoffice-core (LibreOffice)\n"
            "  • DOC antiguo: sudo dnf install antiword (opcional)\n"
            "Abre el archivo en LibreOffice, guárdalo como .docx o .odt y vuelve a intentar."
        )

    return sections_from_plain_text(plain, path.stem)


def read_epub(path: Path) -> list[tuple[str, str]]:
    try:
        import ebooklib
        from ebooklib import epub
        from bs4 import BeautifulSoup
    except ImportError as exc:
        sys.exit(
            "Para EPUB instala dependencias:\n"
            "  pip install --user ebooklib beautifulsoup4\n"
            f"({exc})"
        )

    book = epub.read_epub(str(path))
    sections: list[tuple[str, str]] = []
    for item in book.get_items():
        if item.get_type() != ebooklib.ITEM_DOCUMENT:
            continue
        soup = BeautifulSoup(item.get_content(), "html.parser")
        for tag in soup(["script", "style", "nav"]):
            tag.decompose()
        title = ""
        h = soup.find(["h1", "h2", "h3"])
        if h:
            title = h.get_text(" ", strip=True)
        blocks = []
        for el in soup.find_all(["p", "li", "blockquote", "h1", "h2", "h3", "h4"]):
            t = el.get_text(" ", strip=True)
            if t and len(t) > 2:
                blocks.append(t)
        if not blocks:
            t = soup.get_text("\n", strip=True)
            blocks = [p for p in t.split("\n") if len(p) > 2]
        label = title or f"Sección {len(sections) + 1}"
        body = "\n\n".join(blocks)
        if body.strip():
            sections.append((label, body))
    return sections


def pdf_lines_to_body(text: str) -> str:
    """Une saltos de línea típicos de PDF (corte a mitad de frase)."""
    lines = [ln.strip() for ln in text.split("\n")]
    paragraphs: list[str] = []
    buf = ""
    for line in lines:
        if not line:
            if buf:
                paragraphs.append(buf)
                buf = ""
            continue
        if not buf:
            buf = line
            continue
        ends_sentence = bool(re.search(r'[.!?…:;»"\')]$', buf))
        starts_lower = line[0].islower() if line else False
        if not ends_sentence or starts_lower:
            buf = f"{buf} {line}"
        else:
            paragraphs.append(buf)
            buf = line
    if buf:
        paragraphs.append(buf)
    if paragraphs:
        return "\n\n".join(paragraphs)
    return normalize_text(text)


def clean_ocr_text(text: str) -> str:
    if not text:
        return ""
    # Unir palabras cortadas por guion al final de línea
    text = re.sub(r"(\w+)[-—]\s*\n\s*(\w+)", r"\1\2", text)
    lines = []
    for line in text.splitlines():
        cleaned = line.strip()
        # Filtrar artefactos aislados de ruido en escaneo
        if re.fullmatch(r"[\W_]+", cleaned):
            continue
        lines.append(cleaned)
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def ocr_pixmap(pix, lang: str = "spa+eng") -> str:
    """Ejecuta OCR sobre una imagen renderizada con Tesseract en español e inglés."""
    try:
        import pytesseract
        from PIL import Image

        tessdata = os.path.expanduser("~/.local/share/tessdata")
        if os.path.isdir(tessdata):
            os.environ["TESSDATA_PREFIX"] = tessdata

        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        text = pytesseract.image_to_string(img, lang=lang)
        return clean_ocr_text(text)
    except Exception as e:
        print(f"Aviso OCR: {e}", file=sys.stderr)
        return ""


def extract_tables_from_page(page) -> str:
    """Detecta tablas estructuradas y estadísticas y las convierte en oraciones comprensibles al oído."""
    try:
        tabs = page.find_tables()
        if not tabs or not tabs.tables:
            return ""
        out = []
        for t_idx, tab in enumerate(tabs.tables, 1):
            rows = tab.extract()
            if not rows or len(rows) < 2:
                continue
            headers = [str(h or "").strip() for h in rows[0]]
            out.append(f"\n[Tabla {t_idx}]:")
            for r_idx, row in enumerate(rows[1:], 1):
                cells = []
                for c_idx, val in enumerate(row):
                    val_str = str(val or "").strip()
                    if val_str:
                        col_name = headers[c_idx] if c_idx < len(headers) and headers[c_idx] else f"Columna {c_idx + 1}"
                        cells.append(f"{col_name}: {val_str}")
                if cells:
                    out.append(f"Fila {r_idx}: " + "; ".join(cells) + ".")
        return "\n".join(out)
    except Exception:
        return ""


def extract_figures_and_diagrams(page, existing_text: str) -> str:
    """Detecta imágenes de gráficas o estadísticas y extrae los textos y etiquetas embebidos con OCR."""
    try:
        image_list = page.get_images(full=True)
        if not image_list:
            return ""
        doc = page.parent
        fig_texts = []
        import io
        import pytesseract
        from PIL import Image

        tessdata = os.path.expanduser("~/.local/share/tessdata")
        if os.path.isdir(tessdata):
            os.environ["TESSDATA_PREFIX"] = tessdata

        for img_idx, img_info in enumerate(image_list, 1):
            xref = img_info[0]
            base_img = doc.extract_image(xref)
            if not base_img or not base_img.get("image"):
                continue
            width = base_img.get("width", 0)
            height = base_img.get("height", 0)
            # Descartar iconos o adornos pequeños
            if width < 140 or height < 140:
                continue

            img = Image.open(io.BytesIO(base_img["image"]))
            txt = pytesseract.image_to_string(img, lang="spa+eng").strip()
            cleaned = clean_ocr_text(txt)
            if cleaned and len(cleaned) > 12:
                sample = cleaned[:30]
                if sample not in existing_text:
                    fig_texts.append(f"\n[Datos de Gráfica / Figura {img_idx}]:\n{cleaned}")
        return "\n".join(fig_texts)
    except Exception:
        return ""


def read_pdf_pymupdf(path: Path) -> list[tuple[str, str]] | None:
    try:
        import fitz
    except ImportError:
        return None

    sections: list[tuple[str, str]] = []
    doc = fitz.open(path)
    try:
        if doc.is_encrypted and not doc.authenticate(""):
            sys.exit(
                "Este PDF pide contraseña. Ábrelo en un lector, quita la protección "
                "y guárdalo de nuevo, o exporta a TXT/EPUB."
            )
        for i, page in enumerate(doc):
            vector_text = normalize_text(page.get_text("text"))
            tables_text = extract_tables_from_page(page)

            # Si la página tiene menos de 50 caracteres vectoriales, es un escaneo o foto
            is_scanned = len(vector_text.strip()) < 50
            ocr_text = ""
            if is_scanned:
                # Renderizar página en alta resolución para OCR nítido
                pix = page.get_pixmap(dpi=200)
                ocr_text = ocr_pixmap(pix, lang="spa+eng")

            # Si es texto normal pero contiene gráficas con datos incrustados
            fig_text = ""
            if not is_scanned:
                fig_text = extract_figures_and_diagrams(page, vector_text)

            parts = []
            if ocr_text:
                parts.append(ocr_text)
            elif vector_text:
                parts.append(pdf_lines_to_body(vector_text))

            if tables_text:
                parts.append(tables_text)
            if fig_text:
                parts.append(fig_text)

            body = "\n\n".join(p for p in parts if p.strip()).strip()
            if body:
                label_suffix = " (OCR)" if is_scanned else ""
                sections.append((f"Página {i + 1}{label_suffix}", body))
    finally:
        doc.close()
    return sections if sections else None


def read_pdf_poppler(path: Path) -> list[tuple[str, str]] | None:
    pdftotext = shutil.which("pdftotext")
    if not pdftotext:
        return None
    try:
        proc = subprocess.run(
            [pdftotext, "-layout", str(path), "-"],
            capture_output=True,
            text=True,
            timeout=600,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired):
        return None
    if proc.returncode != 0 or not (proc.stdout or "").strip():
        return None

    raw_pages = proc.stdout.split("\f")
    sections: list[tuple[str, str]] = []
    for i, page in enumerate(raw_pages):
        body = pdf_lines_to_body(normalize_text(page))
        if body.strip():
            sections.append((f"Página {i + 1}", body))
    return sections if sections else None


def read_image_file(path: Path) -> list[tuple[str, str]]:
    """Lee archivos de imagen directa (escaneos, fotos, capturas) usando OCR."""
    try:
        import pytesseract
        from PIL import Image

        tessdata = os.path.expanduser("~/.local/share/tessdata")
        if os.path.isdir(tessdata):
            os.environ["TESSDATA_PREFIX"] = tessdata

        img = Image.open(path)
        text = pytesseract.image_to_string(img, lang="spa+eng")
        cleaned = clean_ocr_text(text)
        if cleaned.strip():
            return [(f"Imagen: {path.name} (OCR)", cleaned)]
    except Exception as e:
        print(f"Error al leer imagen con OCR: {e}", file=sys.stderr)
    return []


def read_pdf(path: Path) -> list[tuple[str, str]]:
    sections = read_pdf_pymupdf(path)
    if sections:
        return sections
    sections = read_pdf_poppler(path)
    if sections:
        return sections

    sys.exit(
        "No se pudo leer el PDF ni extraer texto.\n"
        "Asegúrate de que el archivo no esté corrupto."
    )


def load_sections(path: Path) -> list[tuple[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return read_txt(path)
    if suffix == ".epub":
        return read_epub(path)
    if suffix == ".pdf":
        return read_pdf(path)
    if suffix in (".doc", ".docx", ".docm", ".odt", ".odf", ".rtf", ".wps"):
        return read_office(path)
    if suffix in (".png", ".jpg", ".jpeg", ".webp", ".tiff", ".tif", ".bmp"):
        return read_image_file(path)
    sys.exit(
        f"Formato no soportado: {suffix}. "
        "Usa .pdf, .epub, .txt, .docx, .odt o imágenes (.png, .jpg, .webp)."
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Narra documentos (TXT, EPUB, PDF, Word, ODT, RTF) con RHVoice o Edge TTS."
    )
    parser.add_argument(
        "archivo",
        type=Path,
        help="Ruta al archivo (.txt, .epub, .pdf, .doc, .docx, .odt, .rtf, …)",
    )
    parser.add_argument(
        "--engine",
        choices=("local", "edge", "auto"),
        default="auto",
        help="Motor: local=RHVoice, edge=Microsoft Edge online",
    )
    parser.add_argument(
        "--voice",
        default="",
        help="Voz, p.ej. local:mateo o edge:es-ES-ElviraNeural",
    )
    parser.add_argument("--lang", default="es-ES", help="Código de idioma")
    parser.add_argument("--rate", type=float, default=1.0, help="Velocidad (0.5–1.5)")
    parser.add_argument("--pitch", type=float, default=1.0)
    parser.add_argument("--volume", type=float, default=1.0)
    parser.add_argument(
        "--from-chapter",
        type=int,
        default=1,
        metavar="N",
        help="Empezar en la sección/capítulo/página N (1 = primero; en PDF = página)",
    )
    parser.add_argument(
        "--pause",
        type=float,
        default=0.35,
        help="Pausa en segundos entre fragmentos",
    )
    parser.add_argument(
        "--max-sections",
        type=int,
        default=0,
        help="Solo narrar las primeras N secciones (0 = todas)",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Solo listar secciones y fragmentos, sin hablar",
    )
    return parser.parse_args()


def pick_engine(tts, requested: str) -> str:
    if requested in ("local", "edge"):
        return requested
    engines = tts.detect_engines()
    if "rhvoice" in engines:
        return "local"
    if "edge-tts" in engines:
        return "edge"
    sys.exit("No hay motor de voz. Instala RHVoice o edge-tts (./native/install-rhvoice.sh)")


def pick_voice(tts, engine: str, voice: str, lang: str) -> str:
    if voice:
        return voice
    if engine == "local":
        vid = tts.pick_installed_spanish_voice() if hasattr(tts, "pick_installed_spanish_voice") else "mateo"
        return tts.LOCAL_VOICE_PREFIX + vid
    return tts.pick_neural_voice(lang)


def main() -> int:
    args = parse_args()
    path = args.archivo.expanduser().resolve()
    if not path.is_file():
        sys.exit(f"Archivo no encontrado: {path}")

    tts = load_tts()
    engine = pick_engine(tts, args.engine)
    voice = pick_voice(tts, engine, args.voice, args.lang)

    sections = load_sections(path)
    if not sections:
        sys.exit("No se extrajo texto del archivo.")

    start = max(1, args.from_chapter) - 1
    sections = sections[start:]
    if args.max_sections > 0:
        sections = sections[: args.max_sections]

    print(f"Archivo: {path.name}")
    print(f"Motor: {engine} · Voz: {voice}")
    print(f"Secciones a narrar: {len(sections)}")
    print("Pulsa Ctrl+C para detener.\n")

    total_chunks = 0
    for si, (title, body) in enumerate(sections, start=start + 1):
        print(f"—— {si}. {title[:72]} ——")
        paragraphs = [p.strip() for p in body.split("\n\n") if p.strip()]
        for para in paragraphs:
            for piece in chunk_paragraph(para):
                total_chunks += 1
                preview = piece[:90] + ("…" if len(piece) > 90 else "")
                print(f"  [{total_chunks}] {preview}")
                if args.dry_run:
                    continue
                result = tts.speak(
                    piece,
                    args.lang,
                    args.rate,
                    args.pitch,
                    args.volume,
                    voice,
                    engine,
                )
                if not result.get("ok"):
                    print(f"  Error: {result.get('error', 'desconocido')}", file=sys.stderr)
                    return 1
                if args.pause > 0:
                    time.sleep(args.pause)

    print(f"\nFin ({total_chunks} fragmentos).")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\nDetenido.")
        raise SystemExit(130)
